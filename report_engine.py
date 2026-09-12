from __future__ import annotations

import math
import re
from collections import Counter
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SHEET_NAME = "Order Status-By shipment Date"
ORDER_COUNT = 36
REQUIRED_HEADERS = (
    "Ord Id",
    "Buyer Requested Shipment Date",
    "Prod Id",
    "Prod Name",
    "Pending Prod",
)
REMOVED_HEADERS = {
    "container booking date",
    "shipout date",
    "vsl sailing date",
    "vsl saling date",
    "sale order posted dt. in bc",
    "location name",
}
PLAN_RE = re.compile(r"^planned\s+in\s+(\d+)$", re.IGNORECASE)
CONDITIONS = ("Red", "Black", "Yellow", "Green", "No Fill")
MODE_SHEETS = {
    "all": [
        "Highlighted Report",
        "Red Only",
        "Black Only",
        "Yellow Only",
        "Green Only",
        "Excluded STD Log",
        "Run Summary",
    ],
    "highlighted": ["Highlighted Report"],
    "red": ["Red Only"],
    "black": ["Black Only"],
    "yellow": ["Yellow Only"],
    "green": ["Green Only"],
}

COLORS = {
    "Red": ("C00000", "FFFFFF"),
    "Black": ("000000", "FFFFFF"),
    "Yellow": ("FFD966", "000000"),
    "Green": ("70AD47", "FFFFFF"),
}
DARK = "404040"
BODY_TEXT = "1F1F1F"
MUTED_TEXT = "666666"
LIGHT_LINE = "E7E6E6"


class ReportInputError(ValueError):
    """Raised when the ERP workbook does not meet the locked input contract."""


class ReportValidationError(ValueError):
    """Raised when the generated workbook fails a required reconciliation."""


def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def identifier(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def number(value: Any, *, row: int, header: str) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        raise ReportInputError(
            f"Invalid numeric value at source row {row}, column '{header}': {value!r}"
        )
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        try:
            result = float(str(value).strip().replace(",", ""))
        except ValueError as exc:
            raise ReportInputError(
                f"Invalid numeric value at source row {row}, column '{header}': {value!r}"
            ) from exc
    if not math.isfinite(result):
        raise ReportInputError(
            f"Invalid numeric value at source row {row}, column '{header}': {value!r}"
        )
    return result


def classify(pending: float, planned: list[float]) -> str:
    early = sum(planned[:2])
    late = sum(planned[2:])
    total = early + late
    if pending <= 0:
        return "No Fill"
    if total == 0:
        return "Red"
    if late > 0 and total < pending:
        return "Black"
    if late > 0 and total >= pending:
        return "Yellow"
    if early > 0 and late == 0 and total < pending:
        return "Green"
    return "No Fill"


def _find_header_row(sheet) -> tuple[int, list[Any]]:
    for row_number in range(1, min(sheet.max_row, 60) + 1):
        headers = [sheet.cell(row_number, col).value for col in range(1, sheet.max_column + 1)]
        normalized = {normalize(value) for value in headers if value not in (None, "")}
        if normalize("Ord Id") in normalized and normalize("Prod Name") in normalized:
            return row_number, headers
    raise ReportInputError(f"Could not find the column-header row in sheet '{SHEET_NAME}'.")


def _unique_header_map(headers: list[str]) -> dict[str, int]:
    positions: dict[str, list[int]] = {}
    for index, header in enumerate(headers):
        if header:
            positions.setdefault(normalize(header), []).append(index)
    duplicated_required = [
        header for header in REQUIRED_HEADERS if len(positions.get(normalize(header), [])) > 1
    ]
    if duplicated_required:
        raise ReportInputError(
            "Duplicate required column(s): " + ", ".join(duplicated_required)
        )
    return {key: indexes[0] for key, indexes in positions.items()}


def prepare_report(source: bytes | BinaryIO, source_name: str, order_count: int = ORDER_COUNT) -> dict[str, Any]:
    if not source_name.lower().endswith(".xlsx"):
        raise ReportInputError("Upload an Excel workbook with the .xlsx extension.")
    if order_count < 1:
        raise ReportInputError("The eligible-order count must be at least 1.")

    stream = BytesIO(source) if isinstance(source, bytes) else source
    try:
        workbook = load_workbook(stream, data_only=True, read_only=True)
    except Exception as exc:
        raise ReportInputError("The uploaded file could not be opened as a valid .xlsx workbook.") from exc

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ReportInputError(f"Required sheet not found: '{SHEET_NAME}'.")
        sheet = workbook[SHEET_NAME]
        header_row, raw_headers = _find_header_row(sheet)
        last_header_col = max(
            index for index, value in enumerate(raw_headers) if value not in (None, "")
        )
        headers = [str(value or "").strip() for value in raw_headers[: last_header_col + 1]]
        header_map = _unique_header_map(headers)
        missing = [header for header in REQUIRED_HEADERS if normalize(header) not in header_map]
        if missing:
            raise ReportInputError("Missing required column(s): " + ", ".join(missing))

        plan_source_indexes = [
            index for index, header in enumerate(headers) if PLAN_RE.fullmatch(header)
        ]
        if len(plan_source_indexes) < 2:
            raise ReportInputError(
                "At least two columns named 'Planned in <week number>' are required."
            )
        if plan_source_indexes != list(
            range(plan_source_indexes[0], plan_source_indexes[-1] + 1)
        ):
            raise ReportInputError("Planning-week columns must be contiguous in the source report.")

        order_index = header_map[normalize("Ord Id")]
        prod_id_index = header_map[normalize("Prod Id")]
        prod_name_index = header_map[normalize("Prod Name")]
        pending_index = header_map[normalize("Pending Prod")]

        groups: list[dict[str, Any]] = []
        current_group: dict[str, Any] | None = None
        source_order_seq = 0
        carry = [None] * len(headers)

        for source_row, row in enumerate(
            sheet.iter_rows(
                min_row=header_row + 1,
                max_col=len(headers),
                values_only=True,
            ),
            start=header_row + 1,
        ):
            values = list(row)
            order_value = values[order_index]
            order_id = identifier(order_value)
            starts_new_order = bool(order_id) and (
                current_group is None or order_id != current_group["order_id"]
            )
            if starts_new_order:
                if current_group is not None:
                    groups.append(current_group)
                source_order_seq += 1
                carry = values[:]
                current_group = {
                    "source_order_seq": source_order_seq,
                    "order_id": order_id,
                    "first_source_row": source_row,
                    "last_source_row": source_row,
                    "product_rows": [],
                }
            elif current_group is not None and order_id == current_group["order_id"]:
                for index, value in enumerate(values):
                    if value not in (None, ""):
                        carry[index] = value

            if current_group is None:
                continue
            current_group["last_source_row"] = source_row
            for index in range(prod_id_index):
                if values[index] in (None, "") and carry[index] not in (None, ""):
                    values[index] = carry[index]
                elif values[index] not in (None, ""):
                    carry[index] = values[index]

            fg_id = values[prod_id_index]
            fg_name = str(values[prod_name_index] or "").strip()
            if fg_id in (None, "") or not fg_name:
                continue
            pending = number(
                values[pending_index], row=source_row, header=headers[pending_index]
            )
            planned = [
                number(values[index], row=source_row, header=headers[index])
                for index in plan_source_indexes
            ]
            planned_total = sum(planned)
            current_group["product_rows"].append(
                {
                    "source_row": source_row,
                    "values": values,
                    "fg_id": identifier(fg_id),
                    "fg_name": fg_name,
                    "pending": pending,
                    "planned": planned,
                    "planned_total": planned_total,
                    "unplanned": max(pending - planned_total, 0),
                    "condition": classify(pending, planned),
                }
            )
        if current_group is not None:
            groups.append(current_group)
    finally:
        workbook.close()

    rows: list[dict[str, Any]] = []
    excluded_std: list[dict[str, Any]] = []
    eligible_order_ids: list[str] = []
    output_order_seq = 0
    last_scoped_source_row = None

    for group in groups:
        std_rows = [row for row in group["product_rows"] if "STD" in row["fg_name"].upper()]
        non_std_rows = [
            row for row in group["product_rows"] if "STD" not in row["fg_name"].upper()
        ]
        if not non_std_rows:
            for row in std_rows:
                excluded_std.append(
                    {
                        "output_order_seq": None,
                        "source_order_seq": group["source_order_seq"],
                        "order_id": group["order_id"],
                        "source_row": row["source_row"],
                        "fg_id": row["fg_id"],
                        "fg_name": row["fg_name"],
                        "exclusion_type": "All-STD order removed",
                    }
                )
            continue

        output_order_seq += 1
        eligible_order_ids.append(group["order_id"])
        for row in non_std_rows:
            row = dict(row)
            row["order_seq"] = output_order_seq
            row["source_order_seq"] = group["source_order_seq"]
            rows.append(row)
        for row in std_rows:
            excluded_std.append(
                {
                    "output_order_seq": output_order_seq,
                    "source_order_seq": group["source_order_seq"],
                    "order_id": group["order_id"],
                    "source_row": row["source_row"],
                    "fg_id": row["fg_id"],
                    "fg_name": row["fg_name"],
                    "exclusion_type": "STD FG removed from mixed order",
                }
            )
        last_scoped_source_row = group["last_source_row"]
        if output_order_seq == order_count:
            break

    if output_order_seq < order_count:
        raise ReportInputError(
            f"Only {output_order_seq} eligible orders were found; {order_count} are required."
        )

    kept_source_indexes = [
        index for index, header in enumerate(headers) if normalize(header) not in REMOVED_HEADERS
    ]
    output_headers = ["Order Seq"] + [headers[index] for index in kept_source_indexes]
    normalized_output = {normalize(header) for header in output_headers}
    if normalize("Planned Qty Total") not in normalized_output:
        output_headers.append("Planned Qty Total")
    if normalize("Un Planned") not in normalized_output:
        output_headers.append("Un Planned")
    output_header_map = {normalize(header): index for index, header in enumerate(output_headers)}
    plan_output_indexes = [
        output_header_map[normalize(headers[index])] for index in plan_source_indexes
    ]

    for row in rows:
        output_values = [None] * len(output_headers)
        output_values[0] = row["order_seq"]
        for output_position, source_index in enumerate(kept_source_indexes, start=1):
            output_values[output_position] = row["values"][source_index]
        output_values[output_header_map[normalize("Planned Qty Total")]] = row["planned_total"]
        output_values[output_header_map[normalize("Un Planned")]] = row["unplanned"]
        row["output_values"] = output_values

    counts = Counter(row["condition"] for row in rows)
    for condition in CONDITIONS:
        counts.setdefault(condition, 0)
    all_std_order_ids = sorted(
        {
            item["order_id"]
            for item in excluded_std
            if item["output_order_seq"] is None
        }
    )
    return {
        "source_file": Path(source_name).name,
        "source_sheet": SHEET_NAME,
        "source_header_row": header_row,
        "order_count": order_count,
        "headers": headers,
        "output_headers": output_headers,
        "plan_headers": [headers[index] for index in plan_source_indexes],
        "plan_output_indexes": plan_output_indexes,
        "pending_output_index": output_header_map[normalize("Pending Prod")],
        "total_output_index": output_header_map[normalize("Planned Qty Total")],
        "unplanned_output_index": output_header_map[normalize("Un Planned")],
        "prod_name_output_index": output_header_map[normalize("Prod Name")],
        "rows": rows,
        "excluded_std": excluded_std,
        "condition_counts": dict(counts),
        "eligible_order_ids": eligible_order_ids,
        "all_std_order_ids": all_std_order_ids,
        "last_scoped_source_row": last_scoped_source_row,
    }


def _column_width(header: str) -> float:
    key = normalize(header)
    if key == "order seq":
        return 11
    if key == "prod name":
        return 42
    if key == "buyer name":
        return 28
    if key in {"ord id", "buyer id", "prod id"}:
        return 16
    if "date" in key or key == "plan to ship":
        return 17
    if "contact" in key or "doc no" in key:
        return 20
    if key in {"planned qty total", "un planned"}:
        return 17
    return 15


def _set_sheet_basics(sheet, headers: list[str]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D6"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.outlinePr.summaryBelow = True
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = _column_width(header)


def _write_data_sheet(workbook: Workbook, prepared: dict[str, Any], name: str, condition: str | None) -> None:
    headers = prepared["output_headers"]
    items = prepared["rows"] if condition is None else [
        row for row in prepared["rows"] if row["condition"] == condition
    ]
    sheet = workbook.create_sheet(name)
    _set_sheet_basics(sheet, headers)
    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A1"] = "First 36 eligible orders" if condition is None else f"{condition} planning lines"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=BODY_TEXT)
    scope = (
        f"{prepared['order_count']} eligible orders; STD FG lines excluded."
        if condition is None
        else f"{len(items)} FG {'line' if len(items) == 1 else 'lines'} classified as {condition}."
    )
    sheet["A2"] = f"Source: {prepared['source_file']} | {scope}"
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color=MUTED_TEXT)

    if condition is None:
        legend = [
            (1, "Red", "Completely unplanned"),
            (4, "Black", "Late and partially unplanned"),
            (7, "Yellow", "Late but completely planned"),
            (10, "Green", "Partially planned in the first two planning weeks"),
        ]
        for column, label, description in legend:
            fill_color, font_color = COLORS[label]
            cell = sheet.cell(3, column, label)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(name="Arial", size=9, bold=True, color=font_color)
            sheet.cell(3, column + 1, description).font = Font(
                name="Arial", size=9, color=BODY_TEXT
            )

    white_side = Side(style="thin", color="FFFFFF")
    light_side = Side(style="thin", color=LIGHT_LINE)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(5, column, header)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=white_side, right=white_side, bottom=white_side)
    sheet.row_dimensions[5].height = 42

    start_row = 6
    for row_offset, item in enumerate(items):
        excel_row = start_row + row_offset
        for column, value in enumerate(item["output_values"], start=1):
            cell = sheet.cell(excel_row, column, value)
            cell.font = Font(name="Arial", size=9, color=BODY_TEXT)
            cell.alignment = Alignment(vertical="top")
            cell.border = Border(bottom=light_side)
            header_key = normalize(headers[column - 1])
            if isinstance(value, (datetime, date)):
                cell.number_format = "dd-mmm-yyyy"
            elif header_key in {
                "pending ord qty", "stock in process", "ready stock",
                "stock in process + ready stock", "pending prod", "can pack (in cfc)",
            } or column - 1 in prepared["plan_output_indexes"]:
                cell.number_format = "#,##0.####"
            if header_key in {"prod name", "buyer name"}:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        first_plan = get_column_letter(prepared["plan_output_indexes"][0] + 1)
        last_plan = get_column_letter(prepared["plan_output_indexes"][-1] + 1)
        pending = get_column_letter(prepared["pending_output_index"] + 1)
        total = get_column_letter(prepared["total_output_index"] + 1)
        unplanned = get_column_letter(prepared["unplanned_output_index"] + 1)
        sheet.cell(excel_row, prepared["total_output_index"] + 1).value = (
            f"=SUM({first_plan}{excel_row}:{last_plan}{excel_row})"
        )
        sheet.cell(excel_row, prepared["unplanned_output_index"] + 1).value = (
            f"=MAX({pending}{excel_row}-{total}{excel_row},0)"
        )
        sheet.cell(excel_row, prepared["total_output_index"] + 1).number_format = "#,##0.####"
        sheet.cell(excel_row, prepared["unplanned_output_index"] + 1).number_format = "#,##0.####"

        if item["condition"] in COLORS:
            fill_color, font_color = COLORS[item["condition"]]
            for column in range(
                prepared["plan_output_indexes"][0] + 1,
                prepared["unplanned_output_index"] + 2,
            ):
                cell = sheet.cell(excel_row, column)
                cell.fill = PatternFill("solid", fgColor=fill_color)
                cell.font = Font(name="Arial", size=9, bold=True, color=font_color)
        sheet.row_dimensions[excel_row].height = 30
        sheet.row_dimensions[excel_row].hidden = False
        sheet.row_dimensions[excel_row].outlineLevel = 0
        sheet.row_dimensions[excel_row].collapsed = False

    end_row = max(start_row, start_row + len(items) - 1)
    sheet.auto_filter.ref = f"A5:{last_column}{end_row}"
    sheet.print_title_rows = "5:5"
    sheet.print_area = f"A1:{last_column}{end_row}"
    if not items:
        sheet["A6"] = f"No {condition or 'matching'} FG lines."
        sheet["A6"].font = Font(name="Arial", size=10, italic=True, color=MUTED_TEXT)


def _write_std_log(workbook: Workbook, prepared: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Excluded STD Log")
    headers = [
        "Output Order Seq", "Source Order Seq", "Order ID", "Source Row",
        "FG Item ID", "Prod Name", "Exclusion Type",
    ]
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"
    sheet.merge_cells("A1:G1")
    sheet.merge_cells("A2:G2")
    sheet["A1"] = "Excluded STD finished goods"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=BODY_TEXT)
    sheet["A2"] = f"Source: {prepared['source_file']} | Audit trail for every excluded STD FG line."
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color=MUTED_TEXT)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[4].height = 36
    light_side = Side(style="thin", color=LIGHT_LINE)
    for row_number, item in enumerate(prepared["excluded_std"], start=5):
        values = [
            item["output_order_seq"], item["source_order_seq"], item["order_id"],
            item["source_row"], item["fg_id"], item["fg_name"], item["exclusion_type"],
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Arial", size=9, color=BODY_TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=column in {6, 7})
            cell.border = Border(bottom=light_side)
        sheet.row_dimensions[row_number].height = 30
        sheet.row_dimensions[row_number].hidden = False
    widths = [18, 18, 16, 13, 16, 45, 34]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    end_row = max(5, 4 + len(prepared["excluded_std"]))
    sheet.auto_filter.ref = f"A4:G{end_row}"
    sheet.print_title_rows = "4:4"
    sheet.print_area = f"A1:G{end_row}"


def _write_summary(workbook: Workbook, prepared: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Run Summary")
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    sheet.merge_cells("A2:F2")
    sheet["A1"] = "Order planning run summary"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=BODY_TEXT)
    sheet["A2"] = f"Source: {prepared['source_file']} | Sheet: {prepared['source_sheet']}"
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color=MUTED_TEXT)
    summary = [
        ("Eligible orders", prepared["order_count"]),
        ("Non-STD FG lines", len(prepared["rows"])),
        ("Excluded STD FG lines", len(prepared["excluded_std"])),
        ("All-STD orders removed", len(prepared["all_std_order_ids"])),
        ("Planning columns", ", ".join(prepared["plan_headers"])),
        ("Last scoped source row", prepared["last_scoped_source_row"]),
    ]
    conditions = [
        ("Red", prepared["condition_counts"]["Red"], "Pending production exists and nothing is planned"),
        ("Black", prepared["condition_counts"]["Black"], "Late planning exists and the quantity remains partially unplanned"),
        ("Yellow", prepared["condition_counts"]["Yellow"], "Late planning exists and pending production is completely planned"),
        ("Green", prepared["condition_counts"]["Green"], "Only the first two weeks are planned and the quantity remains partially unplanned"),
        ("No Fill", prepared["condition_counts"]["No Fill"], "No applicable planning exception"),
    ]
    for column, value in enumerate(
        ["Measure", "Value", None, "Condition", "FG lines", "Rule"], start=1
    ):
        sheet.cell(4, column, value)
    for column in (1, 2, 4, 5, 6):
        cell = sheet.cell(4, column)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    for row_index, (label, value) in enumerate(summary, start=5):
        sheet.cell(row_index, 1, label).font = Font(name="Arial", size=10, color=BODY_TEXT)
        sheet.cell(row_index, 2, value).font = Font(name="Arial", size=10, color=BODY_TEXT)
    for row_index, (label, count, rule) in enumerate(conditions, start=5):
        label_cell = sheet.cell(row_index, 4, label)
        label_cell.font = Font(name="Arial", size=10, bold=label in COLORS, color=BODY_TEXT)
        if label in COLORS:
            fill_color, font_color = COLORS[label]
            label_cell.fill = PatternFill("solid", fgColor=fill_color)
            label_cell.font = Font(name="Arial", size=10, bold=True, color=font_color)
        sheet.cell(row_index, 5, count).font = Font(name="Arial", size=10, color=BODY_TEXT)
        rule_cell = sheet.cell(row_index, 6, rule)
        rule_cell.font = Font(name="Arial", size=10, color=BODY_TEXT)
        rule_cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [28, 40, 4, 16, 12, 62]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A4"
    sheet.print_area = "A1:F10"


def generate_workbook(prepared: dict[str, Any], mode: str = "all") -> bytes:
    mode = normalize(mode)
    if mode not in MODE_SHEETS:
        raise ReportValidationError(f"Unsupported report mode: {mode}")
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    if mode in {"all", "highlighted"}:
        _write_data_sheet(workbook, prepared, "Highlighted Report", None)
    for condition in ("Red", "Black", "Yellow", "Green"):
        if mode in {"all", condition.lower()}:
            _write_data_sheet(workbook, prepared, f"{condition} Only", condition)
    if mode == "all":
        _write_std_log(workbook, prepared)
        _write_summary(workbook, prepared)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def validate_generated_workbook(workbook_bytes: bytes, prepared: dict[str, Any], mode: str) -> None:
    mode = normalize(mode)
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=False)
    try:
        expected_sheets = MODE_SHEETS[mode]
        if workbook.sheetnames != expected_sheets:
            raise ReportValidationError(
                f"Generated sheets do not match the selected output: {workbook.sheetnames}"
            )
        data_sheet_names = [
            name for name in expected_sheets if name == "Highlighted Report" or name.endswith(" Only")
        ]
        expected_by_sheet = {
            "Highlighted Report": len(prepared["rows"]),
            "Red Only": prepared["condition_counts"]["Red"],
            "Black Only": prepared["condition_counts"]["Black"],
            "Yellow Only": prepared["condition_counts"]["Yellow"],
            "Green Only": prepared["condition_counts"]["Green"],
        }
        for sheet_name in data_sheet_names:
            sheet = workbook[sheet_name]
            headers = [sheet.cell(5, column).value for column in range(1, sheet.max_column + 1)]
            header_map = {normalize(header): index + 1 for index, header in enumerate(headers)}
            if REMOVED_HEADERS.intersection(header_map):
                raise ReportValidationError(
                    f"Removed columns remain in '{sheet_name}': {REMOVED_HEADERS.intersection(header_map)}"
                )
            expected_rows = expected_by_sheet[sheet_name]
            actual_rows = 0 if expected_rows == 0 else sheet.max_row - 5
            if actual_rows != expected_rows:
                raise ReportValidationError(
                    f"'{sheet_name}' contains {actual_rows} rows; expected {expected_rows}."
                )
            if expected_rows:
                prod_column = header_map[normalize("Prod Name")]
                if any(
                    "STD" in str(sheet.cell(row, prod_column).value or "").upper()
                    for row in range(6, sheet.max_row + 1)
                ):
                    raise ReportValidationError(f"STD FG line found in '{sheet_name}'.")

        if "Highlighted Report" in workbook.sheetnames:
            sheet = workbook["Highlighted Report"]
            sequences = {sheet.cell(row, 1).value for row in range(6, sheet.max_row + 1)}
            if sequences != set(range(1, prepared["order_count"] + 1)):
                raise ReportValidationError("Order Seq does not reconcile to 1–36.")
            headers = [sheet.cell(5, column).value for column in range(1, sheet.max_column + 1)]
            header_map = {normalize(header): index + 1 for index, header in enumerate(headers)}
            first_plan = get_column_letter(prepared["plan_output_indexes"][0] + 1)
            last_plan = get_column_letter(prepared["plan_output_indexes"][-1] + 1)
            pending = get_column_letter(prepared["pending_output_index"] + 1)
            total = get_column_letter(prepared["total_output_index"] + 1)
            for row in (6, sheet.max_row):
                if sheet.cell(row, header_map[normalize("Planned Qty Total")]).value != (
                    f"=SUM({first_plan}{row}:{last_plan}{row})"
                ):
                    raise ReportValidationError("Planned Qty Total formula validation failed.")
                if sheet.cell(row, header_map[normalize("Un Planned")]).value != (
                    f"=MAX({pending}{row}-{total}{row},0)"
                ):
                    raise ReportValidationError("Un Planned formula validation failed.")
        if "Excluded STD Log" in workbook.sheetnames:
            sheet = workbook["Excluded STD Log"]
            actual = max(sheet.max_row - 4, 0)
            if actual != len(prepared["excluded_std"]):
                raise ReportValidationError("Excluded STD Log does not reconcile to the source selection.")
    finally:
        workbook.close()


def preview_rows(prepared: dict[str, Any], limit: int = 200) -> list[dict[str, Any]]:
    headers = prepared["output_headers"]
    wanted = [
        "Order Seq", "Ord Id", "Buyer Name", "Buyer Requested Shipment Date",
        "Prod Id", "Prod Name", "Pending Prod",
    ]
    header_map = {normalize(header): index for index, header in enumerate(headers)}
    result = []
    for item in prepared["rows"][:limit]:
        record = {
            header: item["output_values"][header_map[normalize(header)]]
            for header in wanted
            if normalize(header) in header_map
        }
        record["Planned Qty Total"] = item["planned_total"]
        record["Un Planned"] = item["unplanned"]
        record["Condition"] = item["condition"]
        result.append(record)
    return result


def safe_output_filename(source_name: str, mode: str) -> str:
    stem = Path(source_name).stem
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "", stem).strip()[:80] or "OrderStatus"
    suffix = {
        "all": "Complete",
        "highlighted": "Highlighted",
        "red": "Red_Only",
        "black": "Black_Only",
        "yellow": "Yellow_Only",
        "green": "Green_Only",
    }[normalize(mode)]
    return f"{stem}_First_36_Eligible_Orders_{suffix}.xlsx"
