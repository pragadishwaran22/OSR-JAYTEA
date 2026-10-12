from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from report_engine import (
    ReportInputError,
    generate_workbook,
    prepare_report,
    safe_output_filename,
    validate_generated_workbook,
)


def sample_source() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Order Status-By shipment Date"
    headers = [
        "Ord Id", "Buyer Id", "Buyer Name", "Buyer Requested Shipment Date",
        "Prod Id", "Prod Name", "Pending Prod", "Vsl Saling Date",
        "Planned in 37", "Planned in 38", "Planned in 39",
        "Planned Qty Total", "Un Planned",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(6, column, header)

    source_row = 9
    for order_number in range(1, 41):
        if order_number == 2:
            product_names = ["TEST PRODUCT STD"]
        elif order_number == 1:
            product_names = ["TEST PRODUCT 1", "TEST PRODUCT STD MIXED"]
        elif order_number == 3:
            product_names = ["fg sample product"]
        elif order_number == 4:
            product_names = ["Exempt Product"]
        else:
            product_names = [f"TEST PRODUCT {order_number}"]
        for item_number, product_name in enumerate(product_names):
            pending = 100
            if order_number == 1:
                planned = [0, 0, 0]
            elif order_number == 5:
                planned = [10, 0, 20]
            elif order_number == 6:
                planned = [10, 0, 90]
            elif order_number == 7:
                planned = [30, 0, 0]
            else:
                planned = [100, 0, 0]
            values = [
                f"ORD-{order_number}" if item_number == 0 else None,
                "BUYER-1" if item_number == 0 else None,
                "Sample buyer" if item_number == 0 else None,
                "2026-09-01" if item_number == 0 else None,
                f"FG-{order_number}-{item_number}",
                product_name,
                pending,
                "2026-09-05",
                *planned,
                sum(planned),
                max(pending - sum(planned), 0),
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(source_row, column, value)
            source_row += 1
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def sample_exemptions() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Item_Name"
    sheet["A2"] = "EXEMPT PRODUCT"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class ReportEngineTests(unittest.TestCase):
    def test_buyer_date_orders_whole_orders_before_selecting_36(self):
        workbook = load_workbook(BytesIO(sample_source()))
        sheet = workbook.active
        for row in sheet.iter_rows(min_row=9):
            order_id = row[0].value
            if order_id == "ORD-2":
                row[3].value = "2025-01-01"  # This order is fully excluded.
            elif order_id in {"ORD-39", "ORD-40"}:
                row[3].value = "2026-01-01"  # Equal dates preserve source order.
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        prepared = prepare_report(
            buffer.getvalue(), "daily.xlsx", order_basis="buyer_requested",
            exemption_source=sample_exemptions(),
            exemption_source_name="exemptions.xlsx",
        )
        self.assertEqual(["ORD-39", "ORD-40", "ORD-1"], prepared["eligible_order_ids"][:3])
        self.assertEqual("ORD-37", prepared["eligible_order_ids"][-1])
        self.assertNotIn("ORD-38", prepared["eligible_order_ids"])
        self.assertEqual("buyer_requested", prepared["order_basis"])
        output = generate_workbook(prepared, "all")
        validate_generated_workbook(output, prepared, "all")
        result = load_workbook(BytesIO(output), read_only=True)
        self.assertEqual(
            "Buyer Requested Shipment Date (ascending)",
            result["Run Summary"]["B5"].value,
        )
        result.close()
        self.assertIn(
            "Buyer_Requested_Date",
            safe_output_filename("daily.xlsx", "all", "buyer_requested"),
        )

    def test_complete_workbook_reconciles(self):
        prepared = prepare_report(
            sample_source(), "daily.xlsx", exemption_source=sample_exemptions(),
            exemption_source_name="exemptions.xlsx",
        )
        self.assertEqual(36, len(prepared["eligible_order_ids"]))
        self.assertEqual(36, len(prepared["rows"]))
        self.assertEqual(4, len(prepared["excluded_fgs"]))
        self.assertEqual(["ORD-2", "ORD-3", "ORD-4"], prepared["all_excluded_order_ids"])
        self.assertTrue(any("starts with FG" in row["exclusion_type"] for row in prepared["excluded_fgs"]))
        self.assertTrue(any("Exemption item name" in row["exclusion_type"] for row in prepared["excluded_fgs"]))
        self.assertEqual(1, prepared["condition_counts"]["Red"])
        self.assertEqual(1, prepared["condition_counts"]["Black"])
        self.assertEqual(1, prepared["condition_counts"]["Yellow"])
        self.assertEqual(1, prepared["condition_counts"]["Green"])

        output = generate_workbook(prepared, "all")
        validate_generated_workbook(output, prepared, "all")
        workbook = load_workbook(BytesIO(output), data_only=False)
        self.assertEqual(
            [
                "Highlighted Report", "Red Only", "Black Only", "Yellow Only",
                "Green Only", "Excluded FG Log", "Run Summary",
            ],
            workbook.sheetnames,
        )
        headers = [
            workbook["Highlighted Report"].cell(5, column).value
            for column in range(1, workbook["Highlighted Report"].max_column + 1)
        ]
        self.assertNotIn("Vsl Saling Date", headers)
        workbook.close()

    def test_condition_only_workbook(self):
        prepared = prepare_report(
            sample_source(), "daily.xlsx", exemption_source=sample_exemptions(),
            exemption_source_name="exemptions.xlsx",
        )
        output = generate_workbook(prepared, "red")
        validate_generated_workbook(output, prepared, "red")
        workbook = load_workbook(BytesIO(output), data_only=False)
        self.assertEqual(["Red Only"], workbook.sheetnames)
        self.assertEqual(6, workbook["Red Only"].max_row)
        workbook.close()

    def test_default_exemption_workbook_is_used_without_upload(self):
        prepared = prepare_report(sample_source(), "daily.xlsx")
        self.assertIn(
            "akouna cappuccino 3 in 1 (25gms) 12 strips x 10 sachets with hanger",
            prepared["exemption_names"],
        )

    def test_missing_required_sheet_stops(self):
        workbook = Workbook()
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        with self.assertRaisesRegex(ReportInputError, "Required sheet not found"):
            prepare_report(buffer.getvalue(), "daily.xlsx")

    def test_exemption_list_requires_item_name_header(self):
        workbook = Workbook()
        workbook.active["A1"] = "Product"
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        with self.assertRaisesRegex(ReportInputError, "Item_Name"):
            prepare_report(
                sample_source(), "daily.xlsx", exemption_source=buffer.getvalue(),
                exemption_source_name="exemptions.xlsx",
            )


if __name__ == "__main__":
    unittest.main()
